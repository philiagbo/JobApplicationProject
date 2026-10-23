import contextlib
import io
import logging
import os
import sqlite3
import sys
import tempfile
import time

import docx
from flask.testing import FlaskClient

os.environ["JOBHELPER_DB"] = os.path.join(tempfile.mkdtemp(), "boot.db")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import app  # noqa: E402
import db  # noqa: E402
import sources  # noqa: E402
import watcher  # noqa: E402

sources.REMOTIVE_GAP_SECONDS = 0


class CheckedClient(FlaskClient):
    """Every test checks that failing /api responses carry a JSON error."""

    def open(self, *args, **kwargs):
        resp = super().open(*args, **kwargs)
        if resp.request.path.startswith("/api/") and resp.status_code >= 400:
            body = resp.get_json(silent=True)
            assert isinstance(body, dict) and "error" in body, (resp.request.path, resp.status_code)
        return resp


app.app.test_client_class = CheckedClient

SUMMARY_KEYS = {"salary_text", "salary_annual", "currency", "job_type", "remote",
                "years_experience", "what_youd_do", "what_they_want"}


def item(slug, title, description="", job_type="full_time", salary=""):
    return {"id": slug, "url": "https://remotive.com/remote-jobs/" + slug, "title": title,
            "company_name": "Acme ", "category": "x", "tags": [], "job_type": job_type,
            "publication_date": "2026-10-15", "candidate_required_location": "Worldwide",
            "salary": salary, "description": description}


SHARED = item("shared-1", "Python Developer")
CANNED = {
    "data analyst": [item("da-1", "Data Analyst"), item("om-1", "Office Manager", "we use python daily"), SHARED],
    "python": [SHARED, item("py-2", "Backend Engineer", job_type="contract")],
}

REED_CANNED = {
    "data analyst": [
        {"jobId": 1, "jobUrl": "https://www.reed.co.uk/jobs/data-analyst/1", "jobTitle": "Data Analyst",
         "employerName": "Reed Co ", "locationName": "London", "minimumSalary": 50000, "maximumSalary": 60000,
         "jobDescription": "analyse data"},
        {"jobId": 2, "jobUrl": "https://www.reed.co.uk/jobs/analyst-2/2", "jobTitle": "Analyst",
         "employerName": "Other Co", "locationName": "Leeds", "jobDescription": "no salary shown"},
    ],
}

calls = []
FAIL = set()  # URL substrings that raise OSError


def fake_get_json(url, params, headers=None):
    if any(f in url for f in FAIL):
        raise OSError("down")
    if "remotive.com" in url:
        calls.append(params["search"])
        return {"jobs": CANNED[params["search"]]}
    if "reed.co.uk" in url:
        calls.append("reed:" + params["keywords"])
        return {"results": REED_CANNED[params["keywords"]]}
    raise AssertionError(url)


sources.get_json = fake_get_json


def fresh_client():
    os.environ["JOBHELPER_DB"] = os.path.join(tempfile.mkdtemp(), "t.db")
    db.init_db()
    sources.get_json = fake_get_json  # re-install: another test module may have swapped it
    del calls[:]
    FAIL.clear()
    watcher.STATUS.clear()
    watcher.WAKE.clear()
    return app.app.test_client()


def count_jobs():
    with contextlib.closing(db.connect()) as conn:
        return conn.execute("SELECT COUNT(*) FROM jobs").fetchone()[0]


def test_schema_created():
    fresh_client()
    with contextlib.closing(db.connect()) as conn:
        names = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type = 'table'")}
    assert {"jobs", "answers", "unanswered", "settings"} <= names


def test_settings_defaults():
    c = fresh_client()
    assert c.get("/api/settings").get_json() == db.DEFAULTS
    assert c.get("/api/settings").get_json()["reed_api_key"] == ""


def test_settings_merge():
    c = fresh_client()
    assert c.put("/api/settings", json={"career_fields": [" data analyst ", ""]}).status_code == 200
    r = c.put("/api/settings", json={"min_salary": 40000})
    want = dict(db.DEFAULTS, career_fields=["data analyst"], min_salary=40000)
    assert r.get_json() == want
    assert c.get("/api/settings").get_json() == want


def stored_key():
    with contextlib.closing(db.connect()) as conn:
        return db.get_settings(conn)["reed_api_key"]


def test_settings_key_masked_and_echo_keeps_real_key():
    c = fresh_client()
    r = c.put("/api/settings", json={"reed_api_key": "abc123456789"})
    assert r.get_json()["reed_api_key"] == "\u2022" * 4 + "6789"
    got = c.get("/api/settings").get_json()
    assert got["reed_api_key"].endswith("6789") and "abc123" not in got["reed_api_key"]
    assert b"abc123456789" not in r.data and b"abc123456789" not in c.get("/api/settings").data
    assert c.put("/api/settings", json=dict(got, min_salary=5)).status_code == 200  # form echoes the mask
    assert stored_key() == "abc123456789"
    assert c.put("/api/settings", json={"reed_api_key": "short"}).get_json()["reed_api_key"] == "\u2022" * 8
    assert c.put("/api/settings", json={"reed_api_key": ""}).get_json()["reed_api_key"] == ""
    assert stored_key() == ""


def test_text_settings_are_stripped():
    c = fresh_client()
    c.put("/api/settings", json={"reed_api_key": "  abc123456789\n", "location": " Leeds "})
    assert stored_key() == "abc123456789"
    assert c.get("/api/settings").get_json()["location"] == "Leeds"


def test_settings_rejects_bad_input():
    c = fresh_client()
    bad = [{"nope": 1}, {"min_salary": "40k"}, {"min_salary": True}, {"career_fields": "data"},
           {"career_fields": [1]}, {"api_poll_minutes": 0}, {"min_salary": -1}, {"location": 5},
           [1, 2], {"career_fields": ["ok"], "nope": 1}]
    for body in bad:
        r = c.put("/api/settings", json=body)
        assert r.status_code == 400 and "error" in r.get_json(), body
    assert c.get("/api/settings").get_json() == db.DEFAULTS


def test_fetch_requires_fields():
    c = fresh_client()
    r = c.post("/api/fetch")
    assert r.status_code == 400
    assert r.get_json() == {"error": "Choose at least one career field in settings"}
    assert calls == []


def test_fetch_dedupes():
    c = fresh_client()
    c.put("/api/settings", json={"career_fields": ["data analyst", "python"]})
    r = c.post("/api/fetch")
    assert r.status_code == 200
    assert len(calls) == 2
    assert r.get_json() == {"new": 4, "by_source": {"remotive": 4}, "errors": {}}
    r = c.post("/api/fetch")
    assert r.get_json()["new"] == 0
    assert count_jobs() == 4


def test_jobs_listing():
    c = fresh_client()
    c.put("/api/settings", json={"career_fields": ["data analyst", "python"]})
    c.post("/api/fetch")
    jobs = c.get("/api/jobs").get_json()
    scores = [j["score"] for j in jobs]
    assert scores == sorted(scores, reverse=True)
    for j in jobs:
        assert isinstance(j["pros"], list) and isinstance(j["cons"], list)
        assert isinstance(j["summary"], dict) and set(j["summary"]) == SUMMARY_KEYS
    titles = [j["title"] for j in jobs]
    assert titles.index("Data Analyst") < titles.index("Office Manager")


def test_fetch_source_failure_is_502():
    c = fresh_client()
    c.put("/api/settings", json={"career_fields": ["python"]})
    real = sources.get_json

    def boom(url, params, headers=None):
        raise OSError("down")

    try:
        sources.get_json = boom
        r = c.post("/api/fetch")
        assert r.status_code == 502 and "Remotive" in r.get_json()["error"]
        sources.get_json = lambda url, params, headers=None: {"oops": []}
        r = c.post("/api/fetch")
        assert r.status_code == 502 and "Remotive" in r.get_json()["error"]
    finally:
        sources.get_json = real
    assert count_jobs() == 0


def test_remotive_spacing():
    sources.REMOTIVE_GAP_SECONDS = 0.2
    try:
        start = time.monotonic()
        sources.fetch_remotive(["data analyst", "python"])
        assert time.monotonic() - start >= 0.2
    finally:
        sources.REMOTIVE_GAP_SECONDS = 0


def seeded():
    c = fresh_client()
    c.put("/api/settings", json={"career_fields": ["data analyst", "python"]})
    c.post("/api/fetch")
    return c, {j["title"]: j["id"] for j in c.get("/api/jobs").get_json()}


def sql(query, *args):
    with contextlib.closing(db.connect()) as conn, conn:
        conn.execute(query, args)


def status_changed(job_id):
    with contextlib.closing(db.connect()) as conn:
        return conn.execute("SELECT status_changed_at FROM jobs WHERE id = ?", (job_id,)).fetchone()[0]


def test_dismiss_and_undo():
    c, ids = seeded()
    i = ids["Data Analyst"]
    sql("UPDATE jobs SET status_changed_at = datetime('now', '-1 day') WHERE id = ?", i)
    old = status_changed(i)
    r = c.post("/api/jobs/%d/dismiss" % i)
    assert r.status_code == 200 and r.get_json()["status"] == "dismissed"
    assert status_changed(i) > old
    assert i not in [j["id"] for j in c.get("/api/jobs").get_json()]
    assert i in [j["id"] for j in c.get("/api/jobs?status=dismissed").get_json()]
    sql("UPDATE jobs SET status_changed_at = datetime('now', '-1 day') WHERE id = ?", i)
    old = status_changed(i)
    r = c.post("/api/jobs/%d/undo" % i)
    assert r.get_json()["status"] == "new" and status_changed(i) > old


def test_queue_and_undo():
    c, ids = seeded()
    i = ids["Data Analyst"]
    assert c.post("/api/jobs/%d/queue" % i).get_json()["status"] == "queued"
    assert c.post("/api/jobs/%d/undo" % i).get_json()["status"] == "new"


def test_invalid_transitions_409():
    c, ids = seeded()
    i, k = ids["Data Analyst"], ids["Office Manager"]
    c.post("/api/jobs/%d/queue" % i)
    r = c.post("/api/jobs/%d/dismiss" % i)
    assert r.status_code == 409 and "error" in r.get_json()
    sql("UPDATE jobs SET status = 'applied' WHERE id = ?", k)
    for action in ("dismiss", "queue", "undo"):
        assert c.post("/api/jobs/%d/%s" % (k, action)).status_code == 409


def test_unknown_job_404():
    c = fresh_client()
    for action in ("dismiss", "queue", "undo"):
        r = c.post("/api/jobs/999999/" + action)
        assert r.status_code == 404 and r.get_json() == {"error": "Job not found"}


def test_status_filter():
    c, ids = seeded()
    assert c.get("/api/jobs?status=bogus").status_code == 400
    for t in ("Data Analyst", "Office Manager", "Backend Engineer"):
        c.post("/api/jobs/%d/queue" % ids[t])
    scores = [j["score"] for j in c.get("/api/jobs?status=queued").get_json()]
    assert len(scores) == 3 and scores == sorted(scores, reverse=True)


def empty_metric():
    return {"total": 0, "by_source": {}}


def test_stats_found_window():
    c, ids = seeded()
    assert c.get("/api/stats").get_json()["found"] == {"total": 4, "by_source": {"remotive": 4}}
    sql("UPDATE jobs SET created_at = datetime('now', '-8 days') WHERE id = ?", ids["Data Analyst"])
    s = c.get("/api/stats").get_json()
    assert s["days"] == 7 and s["found"] == {"total": 3, "by_source": {"remotive": 3}}
    for k in ("queued", "dismissed", "applied"):
        assert s[k] == empty_metric()


def test_stats_status_counts():
    c, ids = seeded()
    c.post("/api/jobs/%d/queue" % ids["Data Analyst"])
    c.post("/api/jobs/%d/dismiss" % ids["Office Manager"])
    sql("UPDATE jobs SET status = 'applied' WHERE id = ?", ids["Backend Engineer"])
    one = {"total": 1, "by_source": {"remotive": 1}}
    s = c.get("/api/stats").get_json()
    assert s["queued"] == one and s["dismissed"] == one and s["applied"] == one
    sql("UPDATE jobs SET status_changed_at = datetime('now', '-8 days') WHERE id = ?", ids["Office Manager"])
    assert c.get("/api/stats").get_json()["dismissed"] == empty_metric()


def test_settings_save_rescores_new_only():
    c, ids = seeded()
    c.post("/api/jobs/%d/queue" % ids["Data Analyst"])

    def scores():
        return {j["title"]: j["score"] for s in ("new", "queued") for j in c.get("/api/jobs?status=" + s).get_json()}

    before = scores()
    c.put("/api/settings", json={"career_fields": ["office manager"]})
    after = scores()
    assert after["Office Manager"] > before["Office Manager"]
    assert after["Data Analyst"] == before["Data Analyst"]


def test_excluded_words_never_hide():
    c, _ = seeded()
    c.put("/api/settings", json={"excluded_words": ["analyst"]})
    assert "Data Analyst" in [j["title"] for j in c.get("/api/jobs").get_json()]


def test_unknown_routes_json():
    c = fresh_client()
    assert c.get("/api/nope").status_code == 404
    assert c.post("/api/nope").status_code == 405


def test_non_json_settings_body_400():
    c = fresh_client()
    r = c.put("/api/settings", data="x", content_type="text/plain")
    assert r.status_code == 400 and "error" in r.get_json()


def test_crash_is_500_without_details():
    c = fresh_client()
    real = db.list_jobs

    def boom(conn, status="new"):
        raise RuntimeError("boom")

    app.app.logger.disabled = True  # keep the expected traceback out of test output
    try:
        db.list_jobs = boom
        r = c.get("/api/jobs")
    finally:
        db.list_jobs = real
        app.app.logger.disabled = False
    assert r.status_code == 500 and r.get_json() == {"error": "Internal server error"}
    assert b"boom" not in r.data and b"Traceback" not in r.data


def test_local_only_guard():
    c = fresh_client()
    assert c.get("/api/settings", base_url="http://evil.example").status_code == 403
    assert c.get("/api/settings", headers={"Origin": "https://evil.example"}).status_code == 403
    assert c.get("/api/settings", headers={"Origin": "http://localhost:5173"}).status_code == 200
    assert c.get("/api/settings").status_code == 200


def test_settings_save_wakes_watcher():
    c = fresh_client()
    for body in ({"api_poll_minutes": 5}, {"gmail_poll_minutes": 10}, {"excluded_words": ["x"]}):
        assert c.put("/api/settings", json=body).status_code == 200, body
        assert not watcher.WAKE.is_set(), body
    for body in ({"reed_api_key": "x"}, {"career_fields": ["python"]}, {"min_salary": 1},
                 {"location": "Leeds"}, {"min_salary": 2, "career_fields": ["sql"]}):
        watcher.WAKE.clear()
        assert c.put("/api/settings", json=body).status_code == 200, body
        assert watcher.WAKE.is_set(), body


def test_rejected_settings_save_does_not_wake():
    c = fresh_client()
    assert c.put("/api/settings", json={"career_fields": "data"}).status_code == 400
    assert not watcher.WAKE.is_set()
    assert c.put("/api/settings", data="x", content_type="text/plain").status_code == 400
    assert not watcher.WAKE.is_set()


def test_watcher_status_endpoint_empty():
    c = fresh_client()
    r = c.get("/api/watcher")
    assert r.status_code == 200
    assert r.get_json() == {"running": False, "fetching": False, "sources": {}}
    assert c.get("/api/watcher", headers={"Origin": "https://evil.example"}).status_code == 403


def test_fetch_reed_jobs_scored_and_deduped():
    c = fresh_client()
    c.put("/api/settings", json={"career_fields": ["data analyst"], "reed_api_key": "k"})
    r = c.post("/api/fetch")
    assert r.status_code == 200
    assert r.get_json() == {"new": 5, "by_source": {"reed": 2, "remotive": 3}, "errors": {}}
    reed = [j for j in c.get("/api/jobs").get_json() if j["source"] == "reed"]
    assert len(reed) == 2
    for j in reed:
        assert isinstance(j["pros"], list) and isinstance(j["summary"], dict)
    assert sorted(j["summary"]["salary_annual"] or 0 for j in reed) == [0, 55000]
    r = c.post("/api/fetch")
    assert r.get_json()["by_source"] == {"reed": 0, "remotive": 0}


def test_fetch_partial_failure_is_200():
    c = fresh_client()
    c.put("/api/settings", json={"career_fields": ["data analyst"], "reed_api_key": "SECRET-REED-KEY"})
    FAIL.add("reed.co.uk")
    r = c.post("/api/fetch")
    body = r.get_json()
    assert r.status_code == 200 and body["new"] == 3 and body["by_source"] == {"remotive": 3}
    assert body["errors"]["reed"].startswith("Reed fetch failed")
    w = c.get("/api/watcher")
    srcs = w.get_json()["sources"]
    assert srcs["reed"]["ok"] is False and srcs["reed"]["error"].startswith("Reed fetch failed")
    assert srcs["remotive"]["ok"] is True and srcs["remotive"]["new"] == 3
    assert b"SECRET-REED-KEY" not in r.data and b"SECRET-REED-KEY" not in w.data


def test_fetch_all_sources_failed_is_502():
    c = fresh_client()
    c.put("/api/settings", json={"career_fields": ["data analyst"], "reed_api_key": "k"})
    FAIL.update({"reed.co.uk", "remotive.com"})
    r = c.post("/api/fetch")
    body = r.get_json()
    assert r.status_code == 502 and body["error"].startswith("Reed fetch failed")
    assert set(body["errors"]) == {"reed", "remotive"}
    assert count_jobs() == 0


def test_fetch_busy_is_409():
    c = fresh_client()
    c.put("/api/settings", json={"career_fields": ["data analyst"], "reed_api_key": "k"})
    watcher.FETCH_LOCK.acquire()
    try:
        r = c.post("/api/fetch")
        assert r.status_code == 409 and r.get_json() == {"error": "Fetch already running"}
        assert calls == []
    finally:
        watcher.FETCH_LOCK.release()
    assert c.post("/api/fetch").status_code == 200


def test_fetch_skips_reed_without_key():
    c = fresh_client()
    c.put("/api/settings", json={"career_fields": ["data analyst"]})
    r = c.post("/api/fetch")
    body = r.get_json()
    assert r.status_code == 200 and set(body["by_source"]) == {"remotive"} and body["errors"] == {}
    assert not [x for x in calls if x.startswith("reed:")]
    assert set(c.get("/api/watcher").get_json()["sources"]) == {"remotive"}


# --- saved answers ----------------------------------------------------------

def select(query, *args):
    with contextlib.closing(db.connect()) as conn:
        return conn.execute(query, args).fetchall()


def find(label, options=None):
    with contextlib.closing(db.connect()) as conn:
        return db.find_answer(conn, label, options)


def test_answer_saved_and_matched_end_to_end():
    c = fresh_client()
    r = c.post("/api/answers", json={"question": "Do you require visa sponsorship?", "answer": "No"})
    assert r.status_code in (200, 201)
    assert r.get_json()["question"] == "do you require visa sponsorship" and r.get_json()["answer"] == "No"
    assert {"question": "do you require visa sponsorship", "answer": "No"}.items() <= [
        a for a in c.get("/api/answers").get_json() if a["question"] == "do you require visa sponsorship"][0].items()
    hit = find("Will you require visa sponsorship?", ["Yes", "No"])
    assert hit["option"] == "No" and hit["answer"] == "No" and hit["ratio"] >= 0.85
    r = c.post("/api/answers", json={"question": "Why do you want to work here?",
                                     "answer": "I love Acme's products", "company": "Acme Ltd"})
    assert r.status_code == 400 and r.get_json() == {"error": "Not saved: mentions the company"}
    assert "why do you want to work here" not in [a["question"] for a in c.get("/api/answers").get_json()]


def test_answers_seeded_on_startup():
    c = fresh_client()
    rows = c.get("/api/answers").get_json()
    assert len(rows) == 16 and [r["id"] for r in rows] == sorted(r["id"] for r in rows)
    assert {r["source"] for r in rows} == {"profile"}
    by_q = {r["question"]: r["answer"] for r in rows}
    assert by_q["do you require visa sponsorship"] == ""
    for q in ("what is your gender", "what is your ethnicity", "do you have a disability",
              "what is your sexual orientation", "what is your religion or belief"):
        assert by_q[q] == "Prefer not to say", q
    assert sum(1 for a in by_q.values() if a) == 5
    assert find("Do you require visa sponsorship?") is None
    assert find("What is your gender?", ["Male", "Female", "Prefer not to say"])["option"] == "Prefer not to say"


def test_seed_never_overwrites():
    c = fresh_client()
    seeded = [a for a in c.get("/api/answers").get_json() if a["question"] == "do you require visa sponsorship"][0]
    r = c.post("/api/answers", json={"question": "Do you require visa sponsorship?", "answer": "No"})
    assert r.status_code == 200 and r.get_json()["id"] == seeded["id"]
    db.init_db()
    rows = c.get("/api/answers").get_json()
    assert len(rows) == 16 and [a for a in rows if a["id"] == seeded["id"]][0]["answer"] == "No"


def test_answers_create_then_upsert():
    c = fresh_client()
    r1 = c.post("/api/answers", json={"question": "Do you hold a clean licence?", "answer": "Yes"})
    assert r1.status_code == 201 and r1.get_json()["id"]
    r2 = c.post("/api/answers", json={"question": "do you HOLD a clean licence", "answer": "No"})
    assert r2.status_code == 200 and r2.get_json()["id"] == r1.get_json()["id"] and r2.get_json()["answer"] == "No"
    assert len(select("SELECT id FROM answers WHERE question = 'do you hold a clean licence'")) == 1
    r3 = c.post("/api/answers", json={"question": "do you HOLD a clean licence", "answer": "No"})
    assert r3.status_code == 200 and r3.get_json()["id"] == r1.get_json()["id"]


def test_answers_ordered_by_id():
    c = fresh_client()
    seed = [a["question"] for a in c.get("/api/answers").get_json()]
    assert seed[0] == "what is your address" and len(seed) == 16
    c.post("/api/answers", json={"question": "First new?", "answer": "a"})
    c.post("/api/answers", json={"question": "Second new?", "answer": "b"})
    got = [a["question"] for a in c.get("/api/answers").get_json()]
    assert got == seed + ["first new", "second new"]
    assert [a["question"] for a in c.get("/api/answers").get_json()] == got


def test_answers_edit():
    c = fresh_client()
    a = c.post("/api/answers", json={"question": "Visa?", "answer": "No"}).get_json()
    b = c.post("/api/answers", json={"question": "What is your notice period?", "answer": "1 month"}).get_json()
    r = c.put("/api/answers/%d" % a["id"], json={"answer": "Yes"})
    assert r.status_code == 200 and r.get_json()["answer"] == "Yes" and r.get_json()["question"] == "visa"
    r = c.put("/api/answers/%d" % a["id"], json={"question": "What is your notice period?"})
    assert r.status_code == 409 and r.get_json() == {"error": "An answer for that question already exists"}
    rows = {x["id"]: x for x in c.get("/api/answers").get_json()}
    assert rows[a["id"]]["question"] == "visa" and rows[a["id"]]["answer"] == "Yes"
    assert rows[b["id"]]["question"] == "what is your notice period" and rows[b["id"]]["answer"] == "1 month"
    r = c.put("/api/answers/999999", json={"answer": "x"})
    assert r.status_code == 404 and r.get_json() == {"error": "Answer not found"}
    r = c.put("/api/answers/%d" % a["id"], json={})
    assert r.status_code == 400 and r.get_json() == {"error": "Nothing to change"}


def test_answers_delete():
    c = fresh_client()
    a = c.post("/api/answers", json={"question": "Visa?", "answer": "No"}).get_json()
    r = c.delete("/api/answers/%d" % a["id"])
    assert r.status_code == 200 and r.get_json() == {"ok": True}
    assert a["id"] not in [x["id"] for x in c.get("/api/answers").get_json()]
    r = c.delete("/api/answers/%d" % a["id"])
    assert r.status_code == 404 and r.get_json() == {"error": "Answer not found"}


def test_answers_bad_input():
    c = fresh_client()
    before = len(c.get("/api/answers").get_json())
    bad = [(None, "JSON object"), ([1], "JSON object"), ({"question": "x"}, "answer is required"),
           ({"answer": "a"}, "question is required"), ({"question": "  ", "answer": "a"}, "question is required"),
           ({"question": "???", "answer": "a"}, "question is required"),
           ({"question": "q", "answer": "a", "nope": 1}, "Unknown answer field: nope"),
           ({"question": "q", "answer": 5}, "answer must be text"),
           ({"question": "a" * 501, "answer": "a"}, "question is longer than 500 characters")]
    for body, text in bad:
        r = c.post("/api/answers", data="not json") if body is None else c.post("/api/answers", json=body)
        assert r.status_code == 400 and text in r.get_json()["error"], (body, r.get_json())
    assert len(c.get("/api/answers").get_json()) == before
    r = c.post("/api/answers", json={"question": "Portfolio link?", "answer": ""})
    assert r.status_code == 201 and r.get_json()["answer"] == ""
    assert find("Portfolio link?") is None  # an empty answer never matches


def test_answers_blocked_writes_nothing():
    c = fresh_client()
    n = len(c.get("/api/answers").get_json())
    r = c.post("/api/answers", json={"question": "Tell us more", "answer": "a" * 301})
    assert r.status_code == 400 and r.get_json() == {"error": "Not saved: longer than 300 characters"}
    assert c.post("/api/answers", json={"question": "Tell us more", "answer": "a" * 300}).status_code == 201
    r = c.post("/api/answers", json={"question": "Why do you want to work at Acme?", "answer": "money", "company": "Acme Ltd"})
    assert r.status_code == 400 and r.get_json() == {"error": "Not saved: mentions the company"}
    assert len(c.get("/api/answers").get_json()) == n + 1
    a = c.post("/api/answers", json={"question": "Why us?", "answer": "growth"}).get_json()
    r = c.put("/api/answers/%d" % a["id"], json={"answer": "I admire Acme", "company": "Acme"})
    assert r.status_code == 400 and r.get_json() == {"error": "Not saved: mentions the company"}
    assert select("SELECT answer FROM answers WHERE id = ?", a["id"])[0][0] == "growth"
    assert c.post("/api/answers", json={"question": "Why us?", "answer": "I admire Acme"}).status_code == 200


def test_answer_save_clears_unanswered():
    c = fresh_client()
    sql("INSERT INTO unanswered (question) VALUES ('do you have a full uk driving licence')")
    sql("INSERT INTO unanswered (question) VALUES ('what is your notice period')")
    c.post("/api/answers", json={"question": "Do you have a full UK driving licence?", "answer": "Yes"})
    assert [r[0] for r in select("SELECT question FROM unanswered")] == ["what is your notice period"]


# --- profile and CV upload --------------------------------------------------

def make_pdf(lines):
    esc = lambda s: s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")  # noqa: E731
    stream = "BT /F1 12 Tf 14 TL 72 720 Td " + " ".join("(%s) Tj T*" % esc(ln) for ln in lines) + " ET"
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
            " /Resources << /Font << /F1 5 0 R >> >> >>",
            "<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
            "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += ("%d 0 obj\n%s\nendobj\n" % (i, o)).encode("latin-1")
    x = len(out)
    out += ("xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)).encode()
    for o in offs:
        out += ("%010d 00000 n \n" % o).encode()
    return out + ("trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objs) + 1, x)).encode()


def make_docx(*blocks):
    """A str block is a paragraph; a list of 2-tuples is a 2-column table. Built in the order given."""
    d = docx.Document()
    for block in blocks:
        if isinstance(block, str):
            d.add_paragraph(block)
        else:
            table = d.add_table(rows=len(block), cols=2)
            for row, pair in zip(table.rows, block):
                row.cells[0].text, row.cells[1].text = pair
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


CV_LINES = ["Jane Doe", "jane@example.com | 07700 900123", "Skills", "Excel, SQL (advanced), Python"]


def upload(c, data, filename):
    return c.post("/api/profile/cv", data={"cv": (io.BytesIO(data), filename)}, content_type="multipart/form-data")


def cv_files():
    return sorted(p.name for p in db.cv_dir().iterdir() if p.name.startswith("cv"))


def test_cv_upload_end_to_end():
    c = fresh_client()
    pdf = make_pdf(CV_LINES)
    r = upload(c, pdf, "me.pdf")
    body = r.get_json()
    assert r.status_code == 200 and body["name"] == "Jane Doe" and body["email"] == "jane@example.com"
    assert body["cv_file"] == "cv.pdf"
    got = c.get("/api/profile").get_json()
    assert got["name"] == "Jane Doe" and got["email"] == "jane@example.com"
    assert (db.cv_dir() / "cv.pdf").read_bytes() == pdf


def test_profile_defaults():
    c = fresh_client()
    assert c.get("/api/profile").get_json() == {"name": "", "email": "", "phone": "", "linkedin": "",
                                                "skills": [], "cv_file": ""}
    c.put("/api/profile", json={"name": "Jane"})
    assert "profile" not in c.get("/api/settings").get_json()


def test_profile_put_merge():
    c = fresh_client()
    c.put("/api/profile", json={"name": "Jane Doe"})
    r = c.put("/api/profile", json={"skills": [" SQL ", "sql", "Excel", ""]})
    assert r.status_code == 200 and r.get_json()["name"] == "Jane Doe" and r.get_json()["skills"] == ["SQL", "Excel"]
    assert c.get("/api/profile").get_json()["skills"] == ["SQL", "Excel"]


def test_profile_put_rejects_bad_input():
    c = fresh_client()
    c.put("/api/profile", json={"name": "Jane Doe"})
    before = c.get("/api/profile").get_json()
    assert c.put("/api/profile", data="not json", content_type="text/plain").status_code == 400
    for body in ([1], {"nope": 1}, {"cv_file": "x.pdf"}, {"name": 5}, {"skills": "sql"}, {"skills": [1]}):
        r = c.put("/api/profile", json=body)
        assert r.status_code == 400 and "error" in r.get_json(), body
    assert c.get("/api/profile").get_json() == before


def test_cv_upload_pdf_all_fields():
    c = fresh_client()
    lines = ["Jane Doe", "jane@example.com | 07700 900123", "linkedin.com/in/jane-doe", "Skills",
             "Excel, SQL (advanced), Python"]
    body = upload(c, make_pdf(lines), "cv.pdf").get_json()
    assert (body["name"], body["email"], body["phone"]) == ("Jane Doe", "jane@example.com", "07700 900123")
    assert body["linkedin"] == "https://linkedin.com/in/jane-doe"
    assert body["skills"] == ["Excel", "SQL", "Python"]


def test_cv_upload_docx_with_table():
    c = fresh_client()
    data = make_docx("Sam Patel", [("Phone", "+44 7700 900 123"), ("LinkedIn", "linkedin.com/in/sam-patel")],
                     "Skills", "Excel; Stakeholder management")
    body = upload(c, data, "sam.docx").get_json()
    assert body["name"] == "Sam Patel" and body["phone"] == "+44 7700 900 123"
    assert body["linkedin"] == "https://linkedin.com/in/sam-patel"
    assert body["skills"] == ["Excel", "Stakeholder management"]
    assert body["cv_file"] == "cv.docx"


def test_cv_upload_uses_known_skills_setting():
    c = fresh_client()
    lines = ["Jane Doe", "Skills", "First aid", "Experience", "Safeguarding training for care staff"]
    assert upload(c, make_pdf(lines), "cv.pdf").get_json()["skills"] == ["First aid"]
    c.put("/api/settings", json={"known_skills": ["Safeguarding"]})
    assert upload(c, make_pdf(lines), "cv.pdf").get_json()["skills"] == ["First aid", "Safeguarding"]


def test_cv_upload_keeps_hand_edits():
    c = fresh_client()
    c.put("/api/profile", json={"name": "Janet D", "skills": ["sql"]})
    body = upload(c, make_pdf(CV_LINES), "cv.pdf").get_json()
    assert body["name"] == "Janet D" and body["email"] == "jane@example.com"
    assert body["skills"] == ["sql", "Excel", "Python"]


def test_cv_upload_rejects():
    c = fresh_client()
    c.put("/api/profile", json={"name": "Keep Me"})
    upload(c, make_pdf(CV_LINES), "cv.pdf")
    before = (c.get("/api/profile").get_json(), cv_files(), (db.cv_dir() / "cv.pdf").read_bytes())
    logging.disable(logging.WARNING)  # keep the expected "unreadable CV" warnings out of test output
    try:
        cases = [
            (c.post("/api/profile/cv", data={}, content_type="multipart/form-data"), "No file uploaded"),
            (upload(c, b"hello", "cv.txt"), "Upload a PDF or DOCX file"),
            (upload(c, b"not a pdf", "cv.pdf"), "Couldn't read text from this CV"),
            (upload(c, b"%PDF-" + b"0" * app.MAX_CV_BYTES, "cv.pdf"), "CV is larger than 10 MB"),
            (upload(c, make_pdf([]), "cv.pdf"), "Couldn't read text from this CV"),
            (upload(c, b"PK\x03\x04junk", "cv.docx"), "Couldn't read text from this CV"),
        ]
        # exactly 10 MB passes the size check (it then fails as unreadable, not as too large)
        edge = upload(c, b"%PDF-" + b"0" * (app.MAX_CV_BYTES - 5), "cv.pdf").get_json()
    finally:
        logging.disable(logging.NOTSET)
    for r, error in cases:
        assert r.status_code == 400 and r.get_json() == {"error": error}, (r.status_code, r.get_json(), error)
    assert edge == {"error": "Couldn't read text from this CV"}
    assert (c.get("/api/profile").get_json(), cv_files(), (db.cv_dir() / "cv.pdf").read_bytes()) == before


def test_cv_upload_replaces_file():
    c = fresh_client()
    upload(c, make_pdf(CV_LINES), "a.pdf")
    body = upload(c, make_docx("Sam Patel", "Skills", "Excel"), "b.docx").get_json()
    assert cv_files() == ["cv.docx"] and body["cv_file"] == "cv.docx"
    assert not list(db.cv_dir().glob("*.tmp"))


def test_cv_upload_ignores_client_filename():
    c = fresh_client()
    assert upload(c, make_pdf(CV_LINES), "../../evil.pdf").status_code == 200
    assert cv_files() == ["cv.pdf"]
    assert not list(db.cv_dir().glob("evil*")) and not list(db.cv_dir().parent.glob("evil*"))


if __name__ == "__main__":
    tests = [f for n, f in list(globals().items()) if n.startswith("test_") and callable(f)]
    for t in tests:
        t()
    print("ok: %d tests" % len(tests))
